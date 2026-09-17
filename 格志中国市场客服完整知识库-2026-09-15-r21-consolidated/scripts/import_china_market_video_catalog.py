#!/usr/bin/env python3
"""Normalize the reviewed all-platform tutorial workbook into the current KB.

The workbook is source evidence, not executable instructions. This importer only
copies literal titles/links, applies the confirmed taxonomy, and records the two
review decisions: quarantine Sheet1!E70 and prefer the existing Tmall ribbon URL.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook

try:
    from scripts.package_version import PACKAGE_VERSION
except ModuleNotFoundError:  # Direct execution from scripts/.
    from package_version import PACKAGE_VERSION

SOURCE_SHA256 = "92b8ad5960971caf88e23925e938e4bfa1ab5700eb39d6867812d722d7825127"
PREFERRED_TMALL_RIBBON = (
    "https://cloud.video.taobao.com/vod/"
    "RZwoZfAVTYHTqRO4MjSrj1pp1PAGRMdQI45fvIOwWZk.mp4"
)
PLATFORM_NAMES = {"JD": "JD", "Tmall": "Tmall", "Pinduoduo": "Pinduoduo"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slugify(value: str) -> str:
    value = value.lower().replace("windows", "win")
    value = re.sub(r"[^a-z0-9]+", "_", value).strip("_")
    return value[:90] or "tutorial"


def product_line(source_label: str) -> str | None:
    if source_label == "Smart Attendance Machine":
        return "smart_attendance_machine"
    if source_label == "Needle Printer":
        return "dot_matrix_printer"
    if source_label == "Wifi BT Needle Printer":
        return "wifi_bluetooth_dot_matrix_printer"
    if source_label == "Bluetooth Attendance Machine":
        return "bluetooth_attendance_machine"
    if "Thermal Printer" in source_label:
        return "thermal_printer"
    return None


def operating_system(title_zh: str, title_en: str) -> str | None:
    combined = f"{title_zh} {title_en}".lower()
    if "win10" in combined or "windows 10" in combined:
        return "Windows 10"
    if "win11" in combined or "windows 11" in combined:
        return "Windows 11"
    return None


def keyword_phrases(title_zh: str) -> list[str]:
    phrases = [title_zh]
    cleaned = re.sub(r"^(如何|怎么|怎样)", "", title_zh).strip()
    if cleaned and cleaned != title_zh:
        phrases.append(cleaned)
    replacements = {
        "调节打印速度": ["调整打印速度", "打印速度太慢", "打印速度设置"],
        "打印暂停": ["暂停打印", "打印机暂停", "恢复打印"],
        "无法删除安装包": ["删除不了驱动包", "不能删除驱动安装包", "驱动包无法删除"],
        "定位打印机的IP地址": ["查看打印机IP地址", "查找打印机IP", "打印IP地址"],
        "更换色带": ["色带怎么换", "安装色带", "换色带"],
        "安装驱动": ["驱动安装", "装驱动"],
        "打印位置": ["位置偏移", "打印偏位"],
        "打印不清晰": ["打印模糊", "字迹不清楚"],
        "下载应用程序": ["下载APP", "安装APP"],
        "连接到 Wi-Fi": ["连接WiFi", "手机连打印机WiFi", "iPhone连接打印机"],
        "APP蓝牙连接考勤机": ["蓝牙考勤机APP连接", "APP连接蓝牙考勤机", "蓝牙考勤机设置密码"],
    }
    for needle, aliases in replacements.items():
        if needle.replace(" ", "") in title_zh.replace(" ", ""):
            phrases.extend(aliases)
    return list(dict.fromkeys(phrase for phrase in phrases if phrase))


def build_catalog(workbook: Path) -> dict:
    actual_hash = sha256(workbook)
    if actual_hash != SOURCE_SHA256:
        raise ValueError(f"unexpected source hash: {actual_hash}")

    sheet = load_workbook(workbook, data_only=True, read_only=True)["Sheet1"]
    entries: OrderedDict[tuple[str, str, str | None], dict] = OrderedDict()
    pending_assignments: list[dict] = []
    quarantined_assignments: list[dict] = []

    for row, values in enumerate(
        sheet.iter_rows(min_row=2, max_col=5, values_only=True),
        start=2,
    ):
        platform_raw, source_label, title_zh, title_en, value = values
        if not any(item is not None for item in (platform_raw, source_label, title_zh, title_en, value)):
            continue
        if not all(isinstance(item, str) and item.strip() for item in (platform_raw, source_label, title_zh, title_en)):
            raise ValueError(f"incomplete row {row}")
        platform = PLATFORM_NAMES.get(platform_raw.strip())
        if not platform:
            raise ValueError(f"unsupported platform at row {row}: {platform_raw}")
        cell = f"Sheet1!E{row}"
        value = str(value or "").strip()

        if not value.startswith(("http://", "https://")):
            pending_assignments.append({
                "cell": cell,
                "platform": platform,
                "source_label": source_label.strip(),
                "title_zh": title_zh.strip(),
                "status": "needs_reupload",
                "source_value": value,
            })

        normalized_line = product_line(source_label.strip())
        # Ordinary attendance tutorials keep their richer, later-corrected manifest.
        if normalized_line is None:
            continue

        os_name = operating_system(title_zh, title_en)
        key = (normalized_line, title_zh.strip(), os_name)
        if key not in entries:
            entries[key] = {
                "content_id": f"cloud_{normalized_line}_{slugify(title_en)}",
                "title_zh": title_zh.strip(),
                "title_en": title_en.strip(),
                "product_line": normalized_line,
                "operating_system": os_name,
                "models": [],
                "keywords": keyword_phrases(title_zh.strip()),
                "platform_urls": {},
                "platform_source_cells": {},
                "alternate_platform_urls": {},
                "source_labels": [],
                "analysis_status": "link_verified_only",
            }
        entry = entries[key]
        if source_label.strip() not in entry["source_labels"]:
            entry["source_labels"].append(source_label.strip())

        if cell == "Sheet1!E70":
            quarantined_assignments.append({
                "cell": cell,
                "platform": platform,
                "title_zh": title_zh.strip(),
                "url": value,
                "status": "quarantined_wrong_topic",
                "reason": "Duplicate of Sheet1!E63; frame review supports Win11 speed/settings, not driver-package deletion.",
            })
            continue
        if not value.startswith(("http://", "https://")):
            continue

        if cell == "Sheet1!E113":
            entry["platform_urls"][platform] = PREFERRED_TMALL_RIBBON
            entry["platform_source_cells"][platform] = [
                "customer_service_rag_cards.jsonl:dot_matrix_ribbon_install",
                cell,
            ]
            entry["alternate_platform_urls"][platform] = [value]
        else:
            entry["platform_urls"][platform] = value
            entry["platform_source_cells"][platform] = [cell]

    return {
        "version": PACKAGE_VERSION,
        "catalog_scope": "China-market customer-service tutorial links in the existing KB",
        "source_files": [{
            "file_name": workbook.name,
            "sha256": SOURCE_SHA256,
            "sheet": "Sheet1",
            "imported_on": "2026-09-09",
            "evidence_boundary": "Link and title verified; tutorial steps are not treated as frame-reviewed unless separately documented.",
        }],
        "detailed_manifests": ["attendance_video_materials.json"],
        "platforms": ["JD", "Tmall", "Pinduoduo"],
        "selection_policy": {
            "max_links_per_reply": 1,
            "same_platform_only": True,
            "answer_before_link": True,
            "no_automatic_alternate": True,
            "no_unreviewed_setting_invention": True,
        },
        "entries": list(entries.values()),
        "pending_assignments": pending_assignments,
        "quarantined_assignments": quarantined_assignments,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    catalog = build_catalog(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"topics={len(catalog['entries'])}")
    print(f"active_primary_links={sum(len(item['platform_urls']) for item in catalog['entries'])}")
    print(f"active_alternate_links={sum(sum(len(urls) for urls in item['alternate_platform_urls'].values()) for item in catalog['entries'])}")
    print(f"pending_assignments={len(catalog['pending_assignments'])}")
    print(f"quarantined_assignments={len(catalog['quarantined_assignments'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
